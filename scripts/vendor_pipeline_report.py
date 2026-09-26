#!/usr/bin/env python3
"""Vendor pipeline report in the vendor's own Excel layout -- read-only UAT.

Backlog: docs/Backlog.md "Vendor pipeline report -- Export to Excel for
Admin/GM". Layout: docs/Master List for Q4 2026- Cabio Kerala.xlsx
(SonoScape's 19 columns) + a Sales Owner column. One row per product line
on every open (non-terminal) deal.

    python scripts/vendor_pipeline_report.py --brand SonoScape
    python scripts/vendor_pipeline_report.py --sbu Imaging

--brand needs the Brand/Category/Model catalog (migration 0049+); before
that, use --sbu. Uses the app-role connection from backend/.env.uat in a
read-only session, impersonating an active Admin/GM for RLS (all three
settings, verified). UAT rule (CLAUDE.md): ask Basheer before every run.

The file holds real customer data: it is written to the Desktop (or --out),
never inside the repo.
"""

import argparse
import re
import sys
from copy import copy
from datetime import date
from pathlib import Path

import openpyxl
import psycopg2
import psycopg2.extras
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

sys.stdout.reconfigure(encoding="utf-8")

REPO = Path(__file__).resolve().parent.parent
ENV_FILE = REPO / "backend" / ".env.uat"
TEMPLATE = REPO / "docs" / "Master List for Q4 2026- Cabio Kerala.xlsx"
NCOLS = 20  # the vendor's 19 columns + Sales Owner (vendor asked, 2026-09-26)
COMPETITORS = [  # name shown -> pattern searched in the deal's notes (case-insensitive)
    ("Samsung", r"samsung"), ("GE", r"\bge\b|voluson|logiq|vivid"),
    ("Philips", r"philips|affiniti|epiq"), ("Mindray", r"mindray|resona|\bdc-?\d"),
    ("Siemens", r"siemens|acuson"), ("Canon/Toshiba", r"canon|toshiba|aplio"),
    ("Esaote", r"esaote|mylab"), ("Fujifilm/Sonosite", r"fujifilm|sonosite|arietta"),
    ("Hitachi", r"hitachi|aloka"), ("Vinno", r"vinno"), ("Chison", r"chison"),
    ("Wipro GE", r"wipro"), ("BPL", r"\bbpl\b"), ("Alpinion", r"alpinion"),
]
HINT = PatternFill("solid", fgColor="FFF2CC")

ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
scope = ap.add_mutually_exclusive_group(required=True)
scope.add_argument("--brand", help="only product lines of this brand, e.g. SonoScape")
scope.add_argument("--sbu", help="every line on the SBU's open deals, e.g. Imaging")
ap.add_argument("--out", type=Path, help="output .xlsx (default: Desktop)")
args = ap.parse_args()
label = args.brand or args.sbu
out = args.out or Path.home() / "Desktop" / f"Vendor-Pipeline-{label}-{date.today():%Y-%m-%d}.xlsx"
if REPO in out.resolve().parents:
    sys.exit(f"Refusing to write customer data inside the repo: {out}")

env = {}
for line in ENV_FILE.read_text().splitlines():
    line = line.strip()
    if line and not line.startswith("#") and "=" in line:
        k, v = line.split("=", 1)
        env[k.strip()] = v.strip()

conn = psycopg2.connect(env["DATABASE_URL"])
conn.set_session(readonly=True, autocommit=True)
cur = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
cur.execute("""SELECT up.id, up.sbu_id, up.role_id, up.display_name
  FROM user_profile up JOIN role r ON r.id = up.role_id
  WHERE r.role_name IN ('Admin','General Manager') AND up.is_active = true LIMIT 1""")
a = cur.fetchone()
# is_local=false: this runs under autocommit, so a transaction-local setting
# would vanish before the next statement (see uat_data_quality_check.py).
cur.execute(
    "SELECT set_config('app.current_user_id',%s,false), "
    "set_config('app.current_sbu_id',%s,false), set_config('app.current_role_id',%s,false)",
    (str(a["id"]), str(a["sbu_id"]) if a["sbu_id"] else "", str(a["role_id"])),
)
cur.execute("SELECT cabio_app_uid() uid, cabio_app_role_name() role")
v = cur.fetchone()
if v["uid"] is None or v["role"] not in ("Admin", "General Manager"):
    sys.exit(f"RLS context not set ({v}) -- aborting")
print(f"RLS context: {a['display_name']} ({v['role']})")

if args.brand:
    cur.execute("SELECT to_regclass('public.brand') IS NOT NULL ok")
    if not cur.fetchone()["ok"]:
        sys.exit("No brand table on this database yet (pre-0049) -- use --sbu instead.")
    scope_join = ("JOIN opportunity_item oi ON oi.opportunity_id = o.id "
                  "JOIN product p ON p.id = oi.product_id "
                  "JOIN brand b ON b.id = p.brand_id AND lower(b.name) = lower(%s)")
    scope_arg = args.brand
else:
    scope_join = ("JOIN sbu s ON s.id = o.sbu_id AND lower(s.name) = lower(%s) "
                  "LEFT JOIN opportunity_item oi ON oi.opportunity_id = o.id "
                  "LEFT JOIN product p ON p.id = oi.product_id")
    scope_arg = args.sbu

cur.execute(f"""
SELECT o.id opp_id, o.expected_closure_date, o.competitor_name,
       st.stage_name, os.status_name, os.status_code,
       acc.name acc_name, acc.customer_type, z.name zone_name,
       oi.quantity, oi.line_type, p.name product_name, up.display_name owner
FROM opportunity o
{scope_join}
JOIN opportunity_status os ON os.id = o.status_id AND os.is_terminal = false
JOIN opportunity_stage st ON st.id = o.stage_id
JOIN account acc ON acc.id = o.account_id
LEFT JOIN zone z ON z.id = acc.zone_id
LEFT JOIN user_profile up ON up.id = o.owner_id
ORDER BY z.name, acc.name, o.name, p.name
""", (scope_arg,))
rows = cur.fetchall()
opp_ids = list({r["opp_id"] for r in rows})
print(f"Open deals for {label}: {len(opp_ids)}, lines: {len(rows)}")
if not rows:
    sys.exit("Nothing to report.")


def by_opp(sql):
    cur.execute(sql, (opp_ids,))
    found = {}
    for r in cur.fetchall():
        found.setdefault(r["opp_id"], []).append(r)
    return found


contacts = by_opp("""SELECT os.opportunity_id opp_id, s.name, os.decision_role, os.influence_level
  FROM opportunity_stakeholder os JOIN stakeholder s ON s.id = os.stakeholder_id
  WHERE os.opportunity_id = ANY(%s::uuid[])""")
visits = by_opp("""SELECT opportunity_id opp_id, MAX(activity_date) d FROM activity
  WHERE opportunity_id = ANY(%s::uuid[]) AND activity_type IN ('VISIT','MEETING')
  GROUP BY opportunity_id""")
followups = by_opp("""SELECT a.opportunity_id opp_id, MIN(r.due_date) d FROM reminder r
  JOIN activity a ON a.id = r.activity_id
  WHERE a.opportunity_id = ANY(%s::uuid[]) AND NOT COALESCE(r.is_completed,false)
  GROUP BY a.opportunity_id""")
notes = by_opp("""SELECT opportunity_id opp_id, COALESCE(notes,'') || ' ' || COALESCE(outcome_notes,'') t
  FROM activity WHERE opportunity_id = ANY(%s::uuid[])""")
cur.execute("SELECT DISTINCT opportunity_id FROM opportunity_item "
            "WHERE line_type = 'BUYBACK' AND opportunity_id = ANY(%s::uuid[])", (opp_ids,))
trade_in = {r["opportunity_id"] for r in cur.fetchall()}


def decision_maker(oid):
    cs = contacts.get(oid, [])
    if not cs:
        return None
    def rank(c):
        return (0 if (c["decision_role"] or "").lower().startswith("decision") else 1,
                0 if c["influence_level"] == "HIGH" else 1)
    return sorted(cs, key=rank)[0]["name"]


def competitors(oid, recorded):
    text = " ".join(n["t"] for n in notes.get(oid, [])).lower()
    found = [name for name, pat in COMPETITORS if re.search(pat, text)]
    if recorded:
        found.insert(0, recorded)
    return ", ".join(dict.fromkeys(found)) or None


def day(val):
    return val.date() if hasattr(val, "hour") else val


wb = openpyxl.load_workbook(TEMPLATE)
ws = wb.active
hdr_src = ws.cell(row=1, column=19)
owner_hdr = ws.cell(row=1, column=NCOLS, value="Sales Owner")
owner_hdr.font, owner_hdr.border = copy(hdr_src.font), copy(hdr_src.border)
owner_hdr.alignment, owner_hdr.fill = copy(hdr_src.alignment), copy(hdr_src.fill)
ws.column_dimensions["T"].width = ws.column_dimensions["S"].width or 24
# The template is only pre-formatted to ~row 45, and its empty columns R/S
# carry Excel's default font -- style every data cell from row 2 (column Q
# for the added column) so the look doesn't change down or across the sheet.
style_row = [copy(ws.cell(row=2, column=c if c <= 17 else 17)) for c in range(1, NCOLS + 1)]
row_height = ws.row_dimensions[2].height
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.value = None

sn = 0
for r in rows:
    if r["line_type"] == "BUYBACK":
        continue  # trade-in lines aren't products being sold
    sn += 1
    oid = r["opp_id"]
    ctype = r["customer_type"] or ""
    type1 = ("Government" if ctype == "GOVERNMENT_HOSPITAL"
             else "Check: medical college" if ctype == "MEDICAL_COLLEGE_HOSPITAL"
             else "Private" if ctype else None)
    status = r["stage_name"] if r["status_code"] == "ACTIVE" else f"{r['status_name']} ({r['stage_name']})"
    vals = {
        1: sn, 2: "CABIO Kerala", 3: r["zone_name"], 4: r["acc_name"], 5: type1,
        6: "Replacement" if oid in trade_in else None, 7: None, 8: decision_maker(oid),
        9: None, 10: r["product_name"] or "(no product on deal)", 11: r["quantity"],
        12: None, 13: status, 14: r["expected_closure_date"],
        15: day((visits.get(oid) or [{}])[0].get("d")),
        16: day((followups.get(oid) or [{}])[0].get("d")),
        17: competitors(oid, r["competitor_name"]), 18: None, 19: None, 20: r["owner"],
    }
    for col, val in vals.items():
        cell = ws.cell(row=sn + 1, column=col, value=val)
        t = style_row[col - 1]
        cell.font, cell.border, cell.alignment = copy(t.font), copy(t.border), copy(t.alignment)
        if col in (14, 15, 16) and val:
            cell.number_format = "DD-MMM-YYYY"
    ws.row_dimensions[sn + 1].height = row_height
    if vals[17]:
        ws.cell(row=sn + 1, column=17).fill = HINT
        ws.cell(row=sn + 1, column=17).comment = Comment("From visit notes -- please check", "Cabio Sales OS")
    if type1 == "Check: medical college":
        ws.cell(row=sn + 1, column=5).fill = HINT

scope_text = (f"Open deals with a {args.brand} product, one row per {args.brand} product."
              if args.brand else f"All open {args.sbu} deals, one row per product.")
readme = [
    ("Generated", f"{date.today():%d-%b-%Y} from the Cabio Sales OS (UAT), read-only. {scope_text}"),
    ("Order Status", "The deal's stage. Deals On Hold are included and marked 'On Hold (stage)'."),
    ("City", "The district the hospital is filed under."),
    ("Customer Type I", "Government for government hospitals, Private otherwise. Medical colleges are marked 'Check' (yellow)."),
    ("Customer Type II", "'Replacement' only where the deal includes a trade-in of old equipment. Otherwise blank -- not recorded in the system."),
    ("Area (department)", "Not recorded in the system -- left blank."),
    ("Main Competitor", "Found by searching the deal's visit notes for known competitor names (yellow) -- please check. Not a recorded field."),
    ("Sales Owner", "Extra column (T) the vendor asked for: the salesperson who owns the deal."),
    ("Blank by design", "Purchase Type III, Total Amount, Won with PI No., Lost with Reason."),
]
if args.brand:
    readme.append(("Not included", f"Open deals with no product added can't be matched to {args.brand} and are left out."))
rm = wb.create_sheet("Read me")
rm.column_dimensions["A"].width = 22
rm.column_dimensions["B"].width = 110
for i, (k, t) in enumerate(readme, 1):
    rm.cell(row=i, column=1, value=k).font = Font(bold=True)
    rm.cell(row=i, column=2, value=t)

wb.save(out)
hdr = [c.value for c in ws[1]]
print(f"\nRows written: {sn}\nSaved: {out}\n\nFilled cells per column:")
for col in range(1, NCOLS + 1):
    n = sum(1 for row in ws.iter_rows(min_row=2, max_row=sn + 1) if row[col - 1].value not in (None, ""))
    print(f"  {' '.join(str(hdr[col - 1]).split()):45} {n}/{sn}")
